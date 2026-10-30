"""
Framewise — Streamlit frontend for the AI Video Assistant.

FRONTEND-ONLY REDESIGN. The backend is untouched: the pipeline calls
process_input -> transcribe_all -> generate_title -> summarize ->
extract_action_items / extract_key_decisions / extract_questions ->
build_rag_chain, and chat uses ask_question. Same imports, same arguments,
same order, same environment variables. Nothing in core/ or utils/ changes.

Only functional difference from the previous UI: the Markdown download is
replaced by a Word (.docx) export (needs `python-docx`).
"""
import html
import json
import os
import re
import tempfile
import threading
import time
from datetime import datetime
from io import BytesIO

from dotenv import load_dotenv

load_dotenv()  # before the core imports, so keys from .env are visible to them

import streamlit as st
import streamlit.components.v1 as components
from fpdf import FPDF
from pydub import AudioSegment
from streamlit.runtime.scriptrunner import add_script_run_ctx, get_script_run_ctx

from utils.audio_processor import process_input
from core.transcriber import transcribe_all
from core.summarizer import summarize, generate_title
from core.extractor import extract_action_items, extract_key_decisions, extract_questions
from core.rag_engine import build_rag_chain, ask_question

BRAND = "Framewise"
GITHUB_URL = "https://github.com/Saaksshi18/ai-rag-video-assistant"

st.set_page_config(
    page_title=f"{BRAND} · AI Video Assistant",
    page_icon="▮",
    layout="wide",
    initial_sidebar_state="expanded",
)

SUPPORTED_UPLOAD_TYPES = ["mp4", "mp3", "wav", "m4a"]
SUGGESTED_QUESTIONS = [
    "What decisions were made?",
    "What are my action items?",
    "What concerns were raised?",
]
WAITING_MESSAGES = [
    "Good results take a moment.",
    "Listening closely so nothing gets missed.",
    "Reading between the lines of your transcript.",
    "It hasn't frozen — it's genuinely working.",
    "Almost there — stitching the pieces together.",
]
# (state key, label, what it says while active)
PIPELINE_STEPS = [
    ("audio", "Audio processing", "Fetching and converting the audio"),
    ("transcript", "Transcription", "Turning speech into text — the slowest step"),
    ("summary", "Summary", "Writing the title and executive summary"),
    ("extract", "Insights", "Extracting actions, decisions and open questions"),
    ("rag", "RAG indexing", "Indexing the transcript for chat"),
]
PAGE_TITLES = {"new": "New analysis", "analysis": "Analysis", "history": "Saved & history", "settings": "Settings"}


# ─────────────────────────────────────────────────────────────────────────
# Design system
# ─────────────────────────────────────────────────────────────────────────
CSS = """
<style>
@import url("https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600&family=Manrope:wght@500;600;700&display=swap");

:root {
  --bg:#0e0f11; --sidebar:#101114; --surface:#14161a; --surface-2:#181a1f; --surface-3:#1c1f24;
  --line:#272a30; --line-soft:#202328; --text:#f1f0ed; --text-2:#b1b2b7; --muted:#777a83;
  --violet:#8681dc; --violet-strong:#9691eb; --violet-soft:#201f30; --cyan:#6eb6bf;
  --success:#6eac88; --danger:#d77872; --amber:#c6a46a;
}

html, body, .stApp, [class*="css"] { font-family:"DM Sans", sans-serif; }
.stApp { background:var(--bg); color:var(--text); }
#MainMenu, footer, [data-testid="stDecoration"], [data-testid="stAppDeployButton"], [data-testid="stMainMenu"], [data-testid="stStatusWidget"] { display:none !important; }
[data-testid="stToolbar"] { background:transparent; }
[data-testid="stExpandSidebarButton"], [data-testid="collapsedControl"], [data-testid="stSidebarCollapseButton"] { display:flex !important; visibility:visible !important; color:var(--text-2) !important; }
[data-testid="stHeader"] { background:transparent; height:2.2rem; }
.block-container { max-width:1280px; padding:0.6rem 2.2rem 4rem; }
::selection { background:rgba(134,129,220,.3); }
::-webkit-scrollbar { width:6px; height:6px; }
::-webkit-scrollbar-thumb { background:var(--line); border-radius:3px; }
h1,h2,h3,h4 { font-family:"Manrope", sans-serif; letter-spacing:-.02em; }
hr { border-top:1px solid var(--line-soft) !important; }

/* ── Sidebar ── */
section[data-testid="stSidebar"] { background:var(--sidebar); border-right:1px solid var(--line-soft); }
section[data-testid="stSidebar"] > div { padding-top:0.6rem; }
.logo { display:flex; align-items:center; gap:10px; height:34px; padding:0 4px; margin:4px 0 20px; }
.logo-mark { width:22px; height:22px; display:flex; align-items:flex-end; justify-content:center; gap:2px;
  border:1px solid #3b3e46; border-radius:5px; padding:4px; box-sizing:border-box; }
.logo-mark span { width:2px; background:var(--text); border-radius:1px; }
.logo-mark span:nth-child(1){height:5px} .logo-mark span:nth-child(2){height:11px}
.logo-mark span:nth-child(3){height:8px; background:var(--cyan)}
.logo-name { color:var(--text); font-family:"Manrope",sans-serif; font-weight:600; letter-spacing:-.02em; font-size:1.05rem; }
.side-label { color:#62656d; font-size:10px; font-weight:600; letter-spacing:.06em; text-transform:uppercase; margin:22px 0 6px 4px; }
.side-empty { color:#62656d; font-size:11px; padding:4px; }
section[data-testid="stSidebar"] .stButton > button {
  justify-content:flex-start; background:transparent; border:0; color:var(--muted);
  font-size:12.5px; font-weight:400; min-height:36px; padding:0 10px; border-radius:5px; box-shadow:none; }
section[data-testid="stSidebar"] .stButton > button > div { justify-content:flex-start; width:100%; }
section[data-testid="stSidebar"] .stButton > button:hover { background:var(--surface); color:var(--text-2); }
section[data-testid="stSidebar"] .stButton > button p { text-align:left; overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }
.st-key-nav_new button, .st-key-nav_history button, .st-key-nav_settings button { color:var(--text-2) !important; }

/* ── Header ── */
.topbar { display:flex; align-items:center; justify-content:space-between; padding:0.35rem 0 0.9rem; border-bottom:1px solid var(--line-soft); margin-bottom:1.6rem; }
.breadcrumb { font-size:11.5px; display:flex; gap:8px; align-items:center; }
.breadcrumb span { color:#5e6168; } .breadcrumb b { font-weight:500; color:var(--text-2); }
.topbar-right { display:flex; align-items:center; gap:16px; font-size:11px; color:#73767d; }
.status-dot { display:inline-block; width:6px; height:6px; border-radius:50%; background:var(--success); box-shadow:0 0 0 3px rgba(110,172,136,.08); margin-right:6px; }
.status-dot.warn { background:var(--amber); box-shadow:0 0 0 3px rgba(198,164,106,.1); }
.topbar-right a { color:var(--text-2) !important; text-decoration:none !important; border:1px solid var(--line); padding:5px 11px; border-radius:5px; }
.topbar-right a:hover { background:var(--surface-2); color:var(--text) !important; }

/* ── Type ── */
.context-label, .section-kicker { display:block; color:var(--violet); font-size:11px; font-weight:600; margin-bottom:10px; }
.display-title { font-family:"Manrope",sans-serif; font-size:clamp(34px,4vw,50px); line-height:1.08; letter-spacing:-.045em; font-weight:500; color:var(--text); }
.intro-sub { color:var(--text-2); font-size:15px; line-height:1.6; margin:16px 0 30px; }
.page-title { font-family:"Manrope",sans-serif; font-size:28px; font-weight:500; letter-spacing:-.035em; }
.section-title { font-family:"Manrope",sans-serif; font-size:20px; font-weight:500; letter-spacing:-.025em; color:var(--text); }
.section-title.large { font-size:26px; }
.section-kicker { margin-bottom:6px; }
.section-heading { display:flex; justify-content:space-between; align-items:flex-end; }
.section-heading > span { color:var(--muted); font-size:11px; }

/* ── Buttons / inputs ── */
.stButton > button, .stDownloadButton > button, [data-testid="stPopoverButton"], [data-testid="stBaseButton-secondary"] {
  background:var(--surface-2); color:var(--text-2); border:1px solid var(--line); border-radius:5px;
  font-size:12.5px; font-weight:500; min-height:38px; transition:background .15s, border-color .15s; }
.stButton > button:hover, .stDownloadButton > button:hover, [data-testid="stPopoverButton"]:hover {
  background:var(--surface-3); border-color:#343840; color:var(--text); }
button[kind="primary"], button[data-testid="stBaseButton-primary"] {
  background:var(--violet-strong) !important; border:1px solid var(--violet-strong) !important; color:#111116 !important; font-weight:600 !important; }
button[kind="primary"]:hover, button[data-testid="stBaseButton-primary"]:hover { background:#aaa5f3 !important; border-color:#aaa5f3 !important; }
button:focus-visible { outline:2px solid var(--violet) !important; outline-offset:2px; }
button:disabled { opacity:.42; }
[data-testid="stTextInput"] input, [data-testid="stTextArea"] textarea { background:var(--bg) !important; color:var(--text) !important; }
[data-testid="stTextInput"] div[data-baseweb="input"], [data-testid="stTextArea"] div[data-baseweb="textarea"] {
  background:var(--bg) !important; border:1px solid #353840 !important; border-radius:5px !important; }
[data-testid="stTextInput"] div[data-baseweb="input"], [data-testid="stTextInput"] div[data-baseweb="base-input"] { background:var(--bg) !important; box-shadow:none !important; }
[data-testid="stTextInputRootElement"] { border:1px solid #353840 !important; background:var(--bg) !important; border-radius:5px !important; }
[data-testid="stTextInputRootElement"]:focus-within, [data-testid="stTextInput"] div[data-baseweb="input"]:focus-within { border-color:var(--violet) !important; }
[data-testid="stTextInput"] input { font-size:13.5px; }

/* ── Ingest panel ── */
.st-key-ingest { border:1px solid var(--line) !important; background:var(--surface) !important; border-radius:0 !important; padding:0 !important; gap:0 !important; }
.st-key-ingest > div { gap:0 !important; }
.st-key-method { border-bottom:1px solid var(--line); padding:0 18px; }
.st-key-method div[role="radiogroup"] { gap:26px; }
.st-key-method label { display:flex; padding:0 !important; height:48px; align-items:center; margin:0 !important; position:relative; }
.st-key-method label[data-testid="stRadioOption"] > div > div:first-child { display:none; }
.st-key-method label p { color:var(--muted); font-size:12.5px; }
.st-key-method label:has(input:checked) p { color:var(--text); }
.st-key-method label:has(input:checked)::after { content:""; position:absolute; left:0; right:0; bottom:-1px; height:2px; background:var(--violet); }
.st-key-ingest-body { padding:26px 28px 6px; }
.st-key-lang div[role="radiogroup"] { display:inline-flex; gap:2px; border:1px solid var(--line); padding:2px; border-radius:5px; }
.st-key-lang label { margin:0 !important; padding:0 22px !important; height:30px; border-radius:3px; align-items:center; }
.st-key-lang label[data-testid="stRadioOption"] > div > div:first-child { display:none; }
.st-key-lang label p { font-size:11.5px; color:var(--muted); }
.st-key-lang label:has(input:checked) { background:var(--surface-3); }
.st-key-lang label:has(input:checked) p { color:var(--text); }
.field-label { color:var(--text-2); font-size:11px; margin:18px 0 8px; }
.option-note { color:var(--muted); font-size:11px; line-height:1.5; padding-top:20px; }
.option-note b { color:var(--text-2); font-weight:500; }
.st-key-ingest-footer { border-top:1px solid var(--line); padding:14px 28px; margin-top:20px; align-items:center; }
.ingest-foot-text { color:var(--muted); font-size:11px; padding-top:9px; }
[data-testid="stFileUploaderDropzone"] { background:var(--bg); border:1px dashed #373a42; border-radius:5px; padding:1.6rem; }
[data-testid="stFileUploaderDropzone"]:hover { border-color:#565a65; }
[data-testid="stFileUploader"] small { color:var(--muted); }
.feature-line { display:grid; grid-template-columns:170px 1fr; gap:25px; margin-top:26px; color:#5e6168; font-size:11px; }
.feature-line > div { display:flex; justify-content:space-between; gap:12px; flex-wrap:wrap; }
.feature-line b { color:var(--text-2); font-weight:500; }

/* ── Processing ── */
.processing-header { display:flex; justify-content:space-between; align-items:flex-end; margin-top:6px; }
.processing-header p { color:var(--muted); margin:10px 0 0; }
.process-progress { height:2px; margin:28px 0 34px; background:var(--line); overflow:hidden; }
.process-progress span { display:block; height:100%; background:var(--violet); transition:width .7s ease; }
.steps { padding-left:12px; }
.step { min-height:74px; display:grid; grid-template-columns:38px 1fr auto; gap:14px; align-items:start; position:relative; opacity:.38; }
.step::after { content:""; position:absolute; left:14px; top:30px; bottom:0; width:1px; background:var(--line); }
.step:last-child::after { display:none; }
.step.complete, .step.current { opacity:1; }
.step-number { width:28px; height:28px; border:1px solid var(--line); border-radius:50%; display:grid; place-items:center; color:var(--muted); font-size:10px; background:var(--bg); z-index:1; }
.step.complete .step-number { color:var(--success); border-color:rgba(110,172,136,.35); }
.step.current .step-number { color:var(--violet); border-color:var(--violet); }
.step b, .step small { display:block; } .step b { font-size:13.5px; font-weight:500; margin-top:3px; }
.step small { font-size:11px; color:var(--muted); margin-top:5px; }
.step-pulse { width:7px; height:7px; border-radius:50%; background:var(--violet); animation:pulse 1.5s ease-in-out infinite; margin-top:11px; }
@keyframes pulse { 50% { opacity:.35; } }
.processing-detail { border-left:1px solid var(--line); padding-left:28px; }
.audio-wave { height:64px; display:flex; align-items:center; gap:3px; margin:22px 0 16px; overflow:hidden; }
.audio-wave span { width:2px; min-height:4px; background:#4e5260; border-radius:2px; animation:wave 1.4s ease-in-out infinite alternate; }
.audio-wave span:nth-child(3n) { animation-delay:.3s; background:var(--violet); }
@keyframes wave { to { transform:scaleY(.45); opacity:.5; } }
.wait-timer { font-variant-numeric:tabular-nums; color:var(--violet-strong); font-weight:600; font-size:1.3rem; }
.wait-msg { color:var(--muted); font-size:12px; line-height:1.6; margin-top:6px; }

/* ── Analysis: meeting bar ── */
.meeting-title { font-family:"Manrope",sans-serif; font-weight:600; font-size:21px; letter-spacing:-.02em; }
.meeting-meta { display:flex; flex-wrap:wrap; margin-top:8px; color:var(--muted); font-size:11px; }
.meeting-meta > span + span::before { content:"·"; margin:0 9px; color:#44474e; }
.ready-state { color:var(--success); }
.meeting-rule { border-bottom:1px solid var(--line-soft); margin:18px 0 8px; }

/* ── Analysis: summary ── */
.summary-meta { margin:0; }
.summary-meta > div { padding:0 0 13px; margin-bottom:13px; border-bottom:1px solid var(--line-soft); }
.summary-meta dt { color:var(--muted); font-size:11px; margin-bottom:5px; }
.summary-meta dd { margin:0; color:var(--text-2); font-size:12.5px; word-break:break-word; }
.st-key-summary-copy [data-testid="stMarkdownContainer"] { color:#c7c6c2; font-size:14px; line-height:1.75; }
.st-key-summary-copy strong { color:var(--text); font-weight:500; }
.st-key-summary-copy ul { padding-left:1.2rem; }
.st-key-summary-copy li { margin-bottom:6px; }

/* ── Insights ── */
.insight-heading { display:flex; justify-content:space-between; padding-bottom:11px; border-bottom:1px solid #363940; color:var(--text-2); font-size:11.5px; font-weight:600; }
.insight-heading b { color:var(--muted); font-size:11px; font-weight:500; }
.insight-group.actions .insight-heading { border-bottom-color:rgba(110,182,191,.45); }
.insight-item { display:grid; grid-template-columns:26px 1fr; gap:11px; padding:16px 0; border-bottom:1px solid var(--line-soft); }
.insight-marker { width:24px; height:24px; display:grid; place-items:center; color:var(--muted); background:var(--surface-2); border-radius:3px; font-size:9px; font-weight:600; }
.insight-item b { display:block; font-size:12.5px; line-height:1.45; font-weight:500; color:var(--text); }
.insight-item p { color:var(--text-2); font-size:12px; line-height:1.55; margin:5px 0 0; }
.insight-item small { display:block; margin-top:7px; color:var(--muted); font-size:10.5px; line-height:1.4; }
.insight-empty { color:var(--muted); font-size:12px; padding:16px 0; }
.insight-item strong { font-weight:600; }
.insight-item code { background:var(--surface-3); padding:0 4px; border-radius:3px; font-size:11px; }

/* ── Transcript ── */
.transcript-status { display:flex; gap:20px; margin:8px 0 2px; padding:10px 0; border-top:1px solid var(--line-soft); border-bottom:1px solid var(--line-soft); color:var(--muted); font-size:10.5px; }
.transcript-row { display:grid; grid-template-columns:34px 1fr; gap:12px; padding:18px 0; border-bottom:1px solid var(--line-soft); }
.timestamp { color:var(--cyan); font-variant-numeric:tabular-nums; font-size:11px; padding-top:2px; }
.transcript-row p { margin:0; color:#bbbabe; font-size:13px; line-height:1.72; }
.transcript-row mark { background:rgba(198,164,106,.3); color:var(--text); padding:1px 2px; border-radius:2px; }
.empty-search { padding:40px 0; text-align:center; color:var(--muted); font-size:12px; }
.st-key-transcript-box { border:0 !important; padding:0 !important; }

/* ── Chat panel ── */
div[data-testid="stColumn"]:has(.st-key-chat-panel) { position:sticky; top:2.6rem; align-self:flex-start; }
.st-key-chat-panel { background:#111317 !important; border:1px solid var(--line-soft) !important; border-radius:0 !important; }
.chat-head .section-title { font-size:17px; }
.chat-head p { margin:4px 0 0; color:var(--muted); font-size:11px; }
.chat-context { display:flex; align-items:center; gap:7px; padding:8px 12px; color:var(--text-2); background:#15181c; border:1px solid var(--line-soft); font-size:10.5px; margin:8px 0; }
.chat-context .grow { margin-left:auto; color:var(--muted); }
.st-key-chat-scroll { border:0 !important; padding:0 !important; }
.message-author { display:block; color:var(--violet); font-size:11px; font-weight:600; margin-bottom:6px; }
.message-author.user { color:var(--muted); }
.chat-user { color:var(--text); font-weight:500; font-size:13px; line-height:1.6; margin:0; }
[class*="st-key-msg-"] { padding:14px 0 10px; border-bottom:1px solid var(--line-soft); }
[class*="st-key-msg-"] [data-testid="stMarkdownContainer"] p, [class*="st-key-msg-"] li { font-size:13px; color:#c5c4c6; line-height:1.65; }
.chat-empty { padding:36px 14px 14px; text-align:center; color:var(--muted); }
.chat-empty b { display:block; color:var(--text-2); font-size:13px; margin-bottom:6px; }
.chat-empty p { font-size:11.5px; line-height:1.5; margin:0; }
.chat-disclaimer { text-align:center; color:#555860; font-size:10px; padding:6px 0 0; }
[class*="st-key-sugg-"] button { font-size:11.5px; min-height:32px; justify-content:flex-start; }
[data-testid="stChatInput"], [data-testid="stChatInput"] > div { background:var(--surface-2) !important; border:1px solid #363a43; border-radius:5px; }
[data-testid="stChatInput"] div, [data-testid="stChatInput"] textarea { background:transparent !important; }
[data-testid="stChatInput"] > div { border:0; }
[data-testid="stChatInput"] button { background:transparent !important; color:var(--violet-strong) !important; }
[data-testid="stChatInput"] textarea { color:var(--text); font-size:13px; }
[data-testid="stChatInput"] textarea::placeholder, [data-testid="stTextInput"] input::placeholder { color:#6a6d76 !important; opacity:1; }
[data-testid="stChatInput"]:focus-within { border-color:var(--violet); }

/* ── Popover (export menu) ── */
[data-testid="stPopoverBody"] { background:var(--surface-2); border:1px solid var(--line); border-radius:5px; padding:6px; min-width:230px; }
[data-testid="stPopoverBody"] .stDownloadButton > button, [data-testid="stPopoverBody"] .stButton > button {
  background:transparent; border:0; justify-content:flex-start; color:var(--text-2); font-size:12px; min-height:34px; }
[data-testid="stPopoverBody"] button > div { justify-content:flex-start; width:100%; }
[data-testid="stPopoverBody"] .stDownloadButton > button:hover { background:var(--surface-3); color:var(--text); }

/* ── History / settings ── */
.hist-head { color:var(--muted); font-size:10.5px; padding:10px 0; border-bottom:1px solid var(--line); border-top:1px solid var(--line); }
.hist-title { color:var(--text); font-weight:500; font-size:13px; }
.hist-cell { color:var(--muted); font-size:12px; padding-top:6px; }
.setting-row { display:flex; justify-content:space-between; align-items:center; min-height:66px; border-top:1px solid var(--line); }
.setting-row:last-child { border-bottom:1px solid var(--line); }
.setting-row b { font-size:13px; font-weight:500; } .setting-row p { color:var(--muted); font-size:11px; margin:4px 0 0; }
.pill { font-size:11px; padding:3px 9px; border-radius:20px; border:1px solid var(--line); color:var(--text-2); }
.pill.ok { color:var(--success); border-color:rgba(110,172,136,.35); }
.pill.warn { color:var(--amber); border-color:rgba(198,164,106,.35); }
.settings-title { font-size:12px; font-weight:600; margin:30px 0 10px; }
.danger-title { color:var(--danger); }

.fade-in { animation:fadeIn .35s ease; }
@keyframes fadeIn { from { opacity:0; transform:translateY(4px); } to { opacity:1; transform:none; } }

@media (max-width: 900px) {
  .block-container { padding:0.5rem 1rem 3rem; }
  .feature-line { display:none; }
  div[data-testid="stColumn"]:has(.st-key-chat-panel) { position:static; }
  .processing-detail { border-left:0; padding-left:0; border-top:1px solid var(--line); padding-top:20px; }
}
</style>
"""


def inject_css():
    st.markdown(CSS, unsafe_allow_html=True)


def H(s: str) -> str:
    """Collapse an HTML snippet to one line. Markdown treats indented lines and
    blank lines specially, which would break raw HTML blocks."""
    return re.sub(r"\n\s*", "", s)


def esc(text) -> str:
    return html.escape(text or "")


def format_duration(total_seconds) -> str:
    if not total_seconds or total_seconds <= 0:
        return "—"
    total_seconds = int(total_seconds)
    h, rem = divmod(total_seconds, 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h}h {m}m"
    if m:
        return f"{m}m {s}s"
    return f"{s}s"


def clock(seconds) -> str:
    seconds = int(seconds or 0)
    m, s = divmod(seconds, 60)
    h, m = divmod(m, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


# ─────────────────────────────────────────────────────────────────────────
# Live timer (background thread attached to the script context)
# ─────────────────────────────────────────────────────────────────────────
class LiveTimer:
    """Ticks a placeholder with elapsed time and a rotating reassuring message
    while the pipeline runs. Always stopped in a `finally` by the caller."""

    def __init__(self, placeholder):
        self.placeholder = placeholder
        self.start = time.time()
        self._stop_event = threading.Event()
        self._thread = threading.Thread(target=self._run, daemon=True)
        ctx = get_script_run_ctx()
        if ctx is not None:
            add_script_run_ctx(self._thread, ctx)

    def start_timer(self):
        self._thread.start()

    def _run(self):
        while not self._stop_event.is_set():
            elapsed = int(time.time() - self.start)
            msg = WAITING_MESSAGES[(elapsed // 5) % len(WAITING_MESSAGES)]
            try:
                self.placeholder.markdown(
                    H(f'<div class="wait-timer">{clock(elapsed)}</div><div class="wait-msg">{esc(msg)}</div>'),
                    unsafe_allow_html=True,
                )
            except Exception:
                pass
            self._stop_event.wait(1)

    def stop(self):
        self._stop_event.set()
        self._thread.join(timeout=2)


# ─────────────────────────────────────────────────────────────────────────
# Session state
# ─────────────────────────────────────────────────────────────────────────
DEFAULTS = {
    "result": None,
    "chat_history": [],
    "pipeline_done": False,
    "history": [],
    "page": "new",
    "pending": None,
    "last_error": None,
    "default_language": "English",
}
for _k, _v in DEFAULTS.items():
    if _k not in st.session_state:
        st.session_state[_k] = list(_v) if isinstance(_v, list) else _v


def go(page: str):
    st.session_state.page = page


def reset_to_input():
    st.session_state.result = None
    st.session_state.chat_history = []
    st.session_state.pipeline_done = False
    st.session_state.page = "new"


def open_result(item: dict):
    st.session_state.result = item
    st.session_state.chat_history = []
    st.session_state.pipeline_done = True
    st.session_state.page = "analysis"


# ─────────────────────────────────────────────────────────────────────────
# Text parsing helpers (unchanged logic from the previous UI)
# ─────────────────────────────────────────────────────────────────────────
def split_numbered_items(text: str) -> list:
    if not text or not text.strip():
        return []
    parts = re.split(r"\n(?=\s*\d+[\.\)]\s)", text.strip())
    parts = [p.strip() for p in parts if p.strip()]
    return parts if len(parts) > 1 else [text.strip()]


FIELD_SPLIT = re.compile(r"\b(owner|deadline)\s*[:\-]\s*", re.IGNORECASE)


def parse_action_items(text: str):
    """Parse 'Task ... Owner ... Deadline ...' items. Handles fields on separate
    lines or inline, with or without markdown bold."""
    items = split_numbered_items(text)
    if not items or items[0].lower().startswith("no action items"):
        return []
    rows = []
    for item in items:
        clean = re.sub(r"^\s*\d+[\.\)]\s*", "", item)
        clean = clean.replace("**", "").replace("`", "")
        clean = re.sub(r"^[\s\-\*•]*(task\s*[:\-]\s*)?", "", clean, flags=re.I)
        parts = FIELD_SPLIT.split(clean)
        fields = {"owner": "Not specified", "deadline": "Not specified"}
        for i in range(1, len(parts) - 1, 2):
            fields[parts[i].lower()] = parts[i + 1].strip(" .\n") or "Not specified"
        rows.append({"task": parts[0].strip(" .\n-"), "owner": fields["owner"], "deadline": fields["deadline"]})
    return rows


def inline_md(text: str) -> str:
    s = esc(text)
    s = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s)
    return re.sub(r"`(.+?)`", r"<code>\1</code>", s)


def strip_number(t: str) -> str:
    return re.sub(r"^\s*\d+[.)]\s*", "", t)


def list_items(text: str, empty_prefix: str) -> list:
    items = split_numbered_items(text)
    if len(items) > 1 and not re.match(r"\s*\d+[.)]", items[0]):
        items = items[1:]  # drop a leading un-numbered header line
    if not items or items[0].lower().startswith(empty_prefix.lower()):
        return []
    return [strip_number(i) for i in items]


def title_and_body(text: str):
    """Split one list item into a short bold title and an optional detail line."""
    t = text.strip()
    m = re.match(r"\*\*(.+?)\*\*\s*[:\-–—]?\s*(.*)", t, re.S)
    if m and m.group(1).strip():
        return m.group(1).strip(" :"), m.group(2).strip()
    clean = t.replace("**", "")
    m = re.match(r"(.{12,110}?[.:?!])\s+(.+)", clean, re.S)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return clean, ""


# ─────────────────────────────────────────────────────────────────────────
# Exports
# ─────────────────────────────────────────────────────────────────────────
def build_markdown_report(result: dict) -> str:
    """Used for the 'Copy meeting notes' clipboard button."""
    return (
        f"# {result.get('title', 'Meeting Summary')}\n\n"
        f"**Analyzed:** {result.get('analyzed_at', '')}  \n"
        f"**Source:** {result.get('source', '')}  \n"
        f"**Duration:** {format_duration(result.get('duration_seconds'))}\n\n"
        f"## Executive Summary\n{result.get('summary', '')}\n\n"
        f"## Action Items\n{result.get('action_items', '')}\n\n"
        f"## Key Decisions\n{result.get('key_decisions', '')}\n\n"
        f"## Open Questions\n{result.get('open_questions', '')}\n"
    )


def _xml_safe(text: str) -> str:
    """python-docx raises on control characters that XML cannot hold."""
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text or "")


def _add_runs(paragraph, text: str):
    """Write text into a paragraph, turning **bold** into bold runs."""
    text = _xml_safe(text).replace("`", "")
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if chunk:
            run = paragraph.add_run(chunk)
            run.bold = bool(i % 2)


def build_docx_bytes(result: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2.0)
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(11)

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(0x3B, 0x37, 0x8F)
        return p

    # Title + meta
    t = doc.add_paragraph()
    tr = t.add_run(_xml_safe(result.get("title") or "Meeting Summary"))
    tr.bold = True
    tr.font.size = Pt(22)
    meta = doc.add_paragraph()
    mr = meta.add_run(
        _xml_safe(
            f"Analysed {result.get('analyzed_at', '')}"
            + (f"   |   Duration {format_duration(result.get('duration_seconds'))}" if result.get("duration_seconds") else "")
            + f"   |   Source: {result.get('source', '')}"
        )
    )
    mr.font.size = Pt(9)
    mr.font.color.rgb = RGBColor(0x77, 0x7A, 0x83)

    # Executive summary (bullets / bold aware)
    heading("Executive Summary")
    summary = (result.get("summary") or "").strip()
    if not summary:
        doc.add_paragraph("-")
    for line in summary.splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        indent = len(s) - len(s.lstrip())
        body = s.strip()
        if re.match(r"^[-*•]\s+", body):
            p = doc.add_paragraph(style="List Bullet 2" if indent >= 2 else "List Bullet")
            _add_runs(p, re.sub(r"^[-*•]\s+", "", body))
        elif re.match(r"^#{1,6}\s", body):
            p = doc.add_paragraph()
            r = p.add_run(_xml_safe(re.sub(r"^#{1,6}\s*", "", body)).replace("**", ""))
            r.bold = True
        else:
            p = doc.add_paragraph()
            _add_runs(p, body)

    # Action items table
    heading("Action Items")
    rows = parse_action_items(result.get("action_items", ""))
    if rows:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, ("Task", "Owner", "Due")):
            cell.text = ""
            run = cell.paragraphs[0].add_run(label)
            run.bold = True
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "E6E5F7")
            tcPr.append(shd)
        for r in rows:
            cells = table.add_row().cells
            cells[0].text = _xml_safe(r["task"])
            cells[1].text = _xml_safe(r["owner"])
            cells[2].text = _xml_safe(r["deadline"])
        for row in table.rows:
            row.cells[0].width = Cm(9.0)
            row.cells[1].width = Cm(4.0)
            row.cells[2].width = Cm(3.4)
    else:
        doc.add_paragraph("No action items found.")

    def numbered(title, text, empty):
        heading(title)
        items = list_items(text, empty.split(" found")[0])
        if not items:
            doc.add_paragraph(empty)
            return
        for n, item in enumerate(items, 1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.8)
            p.paragraph_format.space_after = Pt(4)
            p.add_run(f"{n}.  ")
            _add_runs(p, item)

    numbered("Key Decisions", result.get("key_decisions", ""), "No key decisions found.")
    numbered("Open Questions", result.get("open_questions", ""), "No open questions found.")

    # Transcript appendix
    transcript = (result.get("transcript") or "").strip()
    if transcript:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        heading("Transcript")
        for block in transcript_rows(transcript):
            doc.add_paragraph(_xml_safe(block))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_safe(text: str) -> str:
    return (text or "").encode("latin-1", "replace").decode("latin-1")


def build_pdf_bytes(result: dict) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 9, _pdf_safe(result.get("title") or "Meeting Summary"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(130, 130, 130)
    meta = f"Generated {result.get('analyzed_at', '')}   |   Source: {result.get('source', '')}"
    pdf.multi_cell(0, 6, _pdf_safe(meta), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(20, 20, 20)
    pdf.ln(2)
    for heading_text, body in [
        ("Executive Summary", result.get("summary", "")),
        ("Action Items", result.get("action_items", "")),
        ("Key Decisions", result.get("key_decisions", "")),
        ("Open Questions", result.get("open_questions", "")),
    ]:
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 12)
        pdf.multi_cell(0, 7, heading_text, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        pdf.multi_cell(0, 5.5, _pdf_safe(body or "-"), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def safe_bytes(builder, result):
    """Build an export; on any failure return None so the page still renders."""
    try:
        return builder(result)
    except Exception:  # noqa: BLE001
        return None


def save_uploaded_file(uploaded_file) -> str:
    suffix = "." + uploaded_file.name.split(".")[-1].lower()
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    tmp.write(uploaded_file.getbuffer())
    tmp.close()
    return tmp.name


def copy_button(text: str, label: str, key: str):
    """Clipboard copy via a tiny self-contained JS snippet (no extra dependency)."""
    safe_text = json.dumps(text or "")
    safe_key = re.sub(r"[^a-zA-Z0-9_]", "_", key)
    snippet = f"""
        <style>
            :root {{ color-scheme: dark; }}
            html, body {{ margin:0; padding:0; overflow:hidden; background:transparent; }}
        </style>
        <button id="{safe_key}" style="
            width:100%; height:38px; box-sizing:border-box; white-space:nowrap;
            font-family:'DM Sans', 'Source Sans Pro', sans-serif; background:#181a1f; color:#b1b2b7;
            border:1px solid #272a30; border-radius:5px; padding:0 0.9rem;
            font-size:12.5px; font-weight:500; cursor:pointer;">{esc(label)}</button>
        <script>
        const btn_{safe_key} = document.getElementById('{safe_key}');
        btn_{safe_key}.addEventListener('click', function() {{
            navigator.clipboard.writeText({safe_text});
            btn_{safe_key}.innerText = 'Copied';
            setTimeout(() => btn_{safe_key}.innerText = {json.dumps(label)}, 1400);
        }});
        </script>
    """
    if hasattr(st, "iframe"):
        st.iframe(snippet, height=40)
    else:
        components.html(snippet, height=40)


def transcript_rows(transcript: str, per_row: int = 3) -> list:
    """The backend returns one flat string (no timestamps or speakers), so group
    sentences into readable paragraphs instead of inventing either."""
    sentences = [s for s in re.split(r"(?<=[.!?])\s+", (transcript or "").strip()) if s]
    return [" ".join(sentences[i:i + per_row]) for i in range(0, len(sentences), per_row)]


# ─────────────────────────────────────────────────────────────────────────
# Shell: sidebar + top bar
# ─────────────────────────────────────────────────────────────────────────
def nav_button(label: str, key: str, page: str, active: bool):
    if active:
        st.markdown(f"<style>.st-key-{key} button {{ background:var(--surface-2) !important; color:var(--text) !important; }}</style>",
                    unsafe_allow_html=True)
    if st.button(label, key=key, use_container_width=True):
        go(page)
        st.rerun()


def render_sidebar():
    page = st.session_state.page
    with st.sidebar:
        st.markdown(
            H(f'<div class="logo"><span class="logo-mark"><span></span><span></span><span></span></span>'
              f'<span class="logo-name">{BRAND}</span></div>'),
            unsafe_allow_html=True,
        )
        nav_button("＋  New analysis", "nav_new", "new", page == "new")
        nav_button("▤  Saved & history", "nav_history", "history", page == "history")

        st.markdown('<div class="side-label">Recent</div>', unsafe_allow_html=True)
        if st.session_state.history:
            for i, item in enumerate(reversed(st.session_state.history[-6:])):
                label = item["title"][:28] + ("…" if len(item["title"]) > 28 else "")
                is_open = page == "analysis" and st.session_state.result is item
                if st.button(("●  " if is_open else "○  ") + label, key=f"recent_{i}", use_container_width=True):
                    open_result(item)
                    st.rerun()
        else:
            st.markdown('<div class="side-empty">No analyses yet this session.</div>', unsafe_allow_html=True)

        st.markdown('<div class="side-label">Workspace</div>', unsafe_allow_html=True)
        nav_button("⚙  Settings", "nav_settings", "settings", page == "settings")


def render_top_header(page: str):
    key_ok = bool(os.getenv("GROQ_API_KEY"))
    dot = "" if key_ok else " warn"
    status = "Groq key detected" if key_ok else "Groq key missing"
    st.markdown(
        H(f'''<div class="topbar fade-in">
            <div class="breadcrumb"><span>Workspace</span><span>/</span><b>{PAGE_TITLES.get(page, "")}</b></div>
            <div class="topbar-right"><span><i class="status-dot{dot}"></i>{status}</span>
            <a href="{GITHUB_URL}" target="_blank">GitHub</a></div>
        </div>'''),
        unsafe_allow_html=True,
    )


# ─────────────────────────────────────────────────────────────────────────
# Page: New analysis
# ─────────────────────────────────────────────────────────────────────────
def render_new_page():
    err = st.session_state.pop("last_error", None)
    if err:
        st.error(err)

    left, _ = st.columns([5, 1])
    with left:
        st.markdown(
            H('''<div class="fade-in"><span class="context-label">New analysis</span>
            <div class="display-title">Turn hours of video into<br>useful information.</div>
            <div class="intro-sub">Transcribe meetings, extract decisions, and ask questions about your recordings.</div></div>'''),
            unsafe_allow_html=True,
        )

    with st.container(border=True, key="ingest"):
        with st.container(key="method"):
            mode = st.radio("Input method", ["Paste video URL", "Upload video or audio"],
                            horizontal=True, label_visibility="collapsed")
        with st.container(key="ingest-body"):
            source_value, uploaded_file = "", None
            if mode == "Paste video URL":
                source_value = st.text_input("Video URL", placeholder="Paste a YouTube or video URL",
                                             label_visibility="collapsed")
                st.caption("If a YouTube link is blocked (403 / bot check), try again later or upload the file instead.")
            else:
                uploaded_file = st.file_uploader("Upload video or audio", type=SUPPORTED_UPLOAD_TYPES,
                                                 label_visibility="collapsed")
                st.caption("MP4, MP3, WAV or M4A")
            c1, c2 = st.columns([2, 1])
            with c1:
                st.markdown('<div class="field-label">Transcript language</div>', unsafe_allow_html=True)
                with st.container(key="lang"):
                    language_label = st.radio("Language", ["English", "Hinglish"], horizontal=True,
                                              label_visibility="collapsed",
                                              index=0 if st.session_state.default_language == "English" else 1)
            with c2:
                st.markdown(
                    '<div class="option-note">Estimated processing time<br><b>A few minutes, depending on length</b></div>',
                    unsafe_allow_html=True,
                )
        with st.container(key="ingest-footer"):
            f1, f2 = st.columns([3, 1])
            with f1:
                st.markdown('<div class="ingest-foot-text">Results stay in this browser session only.</div>',
                            unsafe_allow_html=True)
            with f2:
                clicked = st.button("Analyse video  →", type="primary", use_container_width=True)

    st.markdown(
        H('''<div class="feature-line"><span>One analysis includes</span>
        <div><b>Accurate transcript</b><b>Executive summary</b><b>Meeting insights</b><b>Grounded Q&amp;A</b></div></div>'''),
        unsafe_allow_html=True,
    )

    if clicked:
        language = "hinglish" if language_label == "Hinglish" else "english"
        if mode == "Paste video URL":
            if not (source_value or "").strip():
                st.error("Please paste a video URL.")
            else:
                st.session_state.pending = {"source": source_value.strip(), "language": language}
                st.rerun()
        else:
            if uploaded_file is None:
                st.error("Please upload a video or audio file.")
            else:
                st.session_state.pending = {"source": save_uploaded_file(uploaded_file), "language": language,
                                            "label": uploaded_file.name}
                st.rerun()


# ─────────────────────────────────────────────────────────────────────────
# Page: Processing (runs the pipeline — identical call order to before)
# ─────────────────────────────────────────────────────────────────────────
def friendly_error(e: Exception) -> str:
    msg = str(e)
    low = msg.lower()
    if "sign in to confirm" in low or "not a bot" in low:
        return (
            "YouTube is asking this server to prove it isn't a bot, so the download was blocked. "
            "This is common for cloud-hosted apps. Use **Upload video or audio** instead, which always works."
        )
    if "403" in msg and ("forbidden" in low or "youtube" in low or "http error 403" in low):
        return (
            "YouTube blocked this server's request (HTTP 403 Forbidden). This happens when YouTube "
            "rate-limits or flags the server's IP address — it isn't something wrong with your link.\n\n"
            "- Wait a minute and try again (the block is often temporary)\n"
            "- Use **Upload video or audio** and upload the file directly"
        )
    if "sarvam_api_key" in low or ("sarvam" in low and "key" in low):
        return "Sarvam API key is missing or invalid. Add `SARVAM_API_KEY` to your environment to use the Hinglish option."
    if "groq_api_key" in low or ("groq" in low and "key" in low):
        return "Groq API key is missing or invalid. Add `GROQ_API_KEY` to your environment."
    return f"Something went wrong: {msg}"


def render_steps(states: dict, steps_ph, progress_ph):
    rows, done, active = [], 0, 0
    for idx, (key, label, active_text) in enumerate(PIPELINE_STEPS, 1):
        state = states.get(key, "pending")
        if state == "done":
            done += 1
        if state == "active":
            active = 1
        cls = {"done": "complete", "active": "current", "pending": ""}[state]
        marker = "✓" if state == "done" else f"0{idx}"
        sub = {"done": "Complete", "active": active_text, "pending": "Waiting"}[state]
        pulse = '<span class="step-pulse"></span>' if state == "active" else "<span></span>"
        rows.append(H(f'<div class="step {cls}"><span class="step-number">{marker}</span>'
                      f'<span><b>{label}</b><small>{sub}</small></span>{pulse}</div>'))
    steps_ph.markdown('<div class="steps">' + "".join(rows) + "</div>", unsafe_allow_html=True)
    pct = int((done + 0.5 * active) / len(PIPELINE_STEPS) * 100)
    progress_ph.markdown(f'<div class="process-progress"><span style="width:{pct}%"></span></div>',
                         unsafe_allow_html=True)


def run_pipeline(source: str, language: str, label: str):
    st.session_state.result = None
    st.session_state.chat_history = []
    start_time = datetime.now()

    st.markdown('<span class="context-label">Processing recording</span>', unsafe_allow_html=True)
    shown_label = "Analysing your video" if label.startswith("http") else label
    st.markdown(
        H(f'<div class="processing-header"><div><div class="page-title">{esc(shown_label)}</div>'
          f'<p>We’re turning your recording into a searchable workspace.</p></div></div>'),
        unsafe_allow_html=True,
    )
    progress_ph = st.empty()
    left, right = st.columns([2.2, 1], gap="large")
    with left:
        steps_ph = st.empty()
    with right:
        bars = "".join(f'<span style="height:{12 + (i * 13) % 34}px"></span>' for i in range(30))
        st.markdown(
            f'<div class="processing-detail"><div class="side-label" style="margin:0">Elapsed</div>'
            f'<div class="audio-wave">{bars}</div></div>',
            unsafe_allow_html=True,
        )
        timer_ph = st.empty()

    states = {key: "pending" for key, _, _ in PIPELINE_STEPS}
    render_steps(states, steps_ph, progress_ph)

    def mark(key, state):
        states[key] = state
        render_steps(states, steps_ph, progress_ph)

    timer = LiveTimer(timer_ph)
    timer.start_timer()
    try:
        mark("audio", "active")
        chunks = process_input(source)
        mark("audio", "done")

        # Real duration from the chunk files audio_processor already produced.
        duration_seconds = 0.0
        try:
            for c in chunks:
                duration_seconds += len(AudioSegment.from_wav(c)) / 1000.0
        except Exception:
            duration_seconds = 0.0

        mark("transcript", "active")
        transcript = transcribe_all(chunks, language)
        mark("transcript", "done")

        mark("summary", "active")
        title = generate_title(transcript)
        summary = summarize(transcript)
        mark("summary", "done")

        mark("extract", "active")
        action_items = extract_action_items(transcript)
        decisions = extract_key_decisions(transcript)
        questions = extract_questions(transcript)
        mark("extract", "done")

        mark("rag", "active")
        rag_chain = build_rag_chain(transcript)
        mark("rag", "done")
    except Exception as e:  # noqa: BLE001
        timer.stop()
        st.session_state.last_error = friendly_error(e)
        st.session_state.page = "new"
        st.rerun()
    finally:
        timer.stop()

    elapsed = (datetime.now() - start_time).total_seconds()
    result = {
        "title": title,
        "transcript": transcript,
        "summary": summary,
        "action_items": action_items,
        "key_decisions": decisions,
        "open_questions": questions,
        "rag_chain": rag_chain,
        "source": source if source.startswith("http") else label,
        "language": language,
        "duration_seconds": duration_seconds,
        "analyzed_at": datetime.now().strftime("%d %b %Y, %H:%M"),
        "processing_seconds": elapsed,
    }
    st.session_state.result = result
    st.session_state.history.append(result)
    st.session_state.pipeline_done = True
    st.session_state.page = "analysis"
    st.rerun()


# ─────────────────────────────────────────────────────────────────────────
# Page: Analysis
# ─────────────────────────────────────────────────────────────────────────
def render_meeting_bar(result: dict):
    left, right = st.columns([5, 3], gap="medium")
    with left:
        src = result.get("source", "")
        dur = format_duration(result.get("duration_seconds"))
        dur_html = f"<span>{dur}</span>" if dur != "—" else ""
        src_label = "YouTube" if "youtu" in src else ("Link" if src.startswith("http") else "Upload")
        st.markdown(
            H(f'''<div class="meeting-title">{esc(result["title"])}</div>
            <div class="meeting-meta"><span>{esc(result.get("analyzed_at", ""))}</span>
            {dur_html}<span>{src_label}</span>
            <span class="ready-state">✓ Analysis complete</span></div>'''),
            unsafe_allow_html=True,
        )
    with right:
        b1, b2 = st.columns(2)
        with b1:
            with st.popover("Export", use_container_width=True):
                docx_bytes = safe_bytes(build_docx_bytes, result)
                if docx_bytes:
                    st.download_button("Word document (.docx)", data=docx_bytes, key="dl_docx",
                                       file_name="meeting-report.docx",
                                       mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                       use_container_width=True)
                else:
                    st.button("Word export unavailable (install python-docx)", disabled=True, key="dl_docx_off",
                              use_container_width=True)
                pdf_bytes = safe_bytes(build_pdf_bytes, result)
                if pdf_bytes:
                    st.download_button("PDF", data=pdf_bytes, key="dl_pdf", file_name="meeting-report.pdf",
                                       mime="application/pdf", use_container_width=True)
                st.download_button("Transcript (.txt)", data=result.get("transcript", ""), key="dl_txt",
                                   file_name="transcript.txt", mime="text/plain", use_container_width=True)
                copy_button(build_markdown_report(result), "Copy meeting notes", key="copy_notes")
        with b2:
            if st.button("＋  New analysis", type="primary", use_container_width=True, key="bar_new"):
                reset_to_input()
                st.rerun()
    st.markdown('<div class="meeting-rule"></div>', unsafe_allow_html=True)


def render_summary_section(result: dict):
    st.markdown(
        '<span class="section-kicker">Overview</span><div class="section-title large">Executive Summary</div>',
        unsafe_allow_html=True,
    )
    st.write("")
    c1, c2 = st.columns([1.5, 0.6], gap="large")
    with c1:
        with st.container(key="summary-copy"):
            st.markdown(result.get("summary", "") or "—")
    with c2:
        engine = "Whisper (English)" if result.get("language") == "english" else "Sarvam AI (Hinglish → English)"
        src = result.get("source", "")
        st.markdown(
            H(f'''<dl class="summary-meta">
            <div><dt>Duration</dt><dd>{format_duration(result.get("duration_seconds"))}</dd></div>
            <div><dt>Engine</dt><dd>{esc(engine)}</dd></div>
            <div><dt>Processed in</dt><dd>{format_duration(result.get("processing_seconds"))}</dd></div>
            <div><dt>Source</dt><dd>{esc(src if len(src) < 60 else src[:57] + "…")}</dd></div></dl>'''),
            unsafe_allow_html=True,
        )


def insight_items_html(entries: list, prefix: str, empty: str) -> str:
    if not entries:
        return f'<div class="insight-empty">{esc(empty)}</div>'
    out = []
    for n, (title, body, meta) in enumerate(entries, 1):
        body_html = f"<p>{inline_md(body)}</p>" if body else ""
        meta_html = f"<small>{esc(meta)}</small>" if meta else ""
        out.append(f'<div class="insight-item"><span class="insight-marker">{prefix}{n}</span>'
                   f'<div><b>{inline_md(title)}</b>{body_html}{meta_html}</div></div>')
    return "".join(out)


def render_insights_section(result: dict):
    # Action items
    actions = []
    for r in parse_action_items(result.get("action_items", "")):
        owner = r["owner"] if r["owner"].lower() != "not specified" else "Unassigned"
        due = f"Due {r['deadline']}" if r["deadline"].lower() != "not specified" else "No deadline"
        actions.append((r["task"], "", f"{owner} · {due}"))
    decisions = [(*title_and_body(t), "") for t in list_items(result.get("key_decisions", ""), "No key decisions")]
    questions = [(*title_and_body(t), "") for t in list_items(result.get("open_questions", ""), "No open questions")]
    total = len(actions) + len(decisions) + len(questions)

    st.markdown(
        H(f'<div class="section-heading"><div><span class="section-kicker">Structured notes</span>'
          f'<div class="section-title">Key Insights</div></div><span>{total} items extracted</span></div>'),
        unsafe_allow_html=True,
    )
    st.write("")
    c1, c2, c3 = st.columns(3, gap="large")
    cols = [
        (c1, "Action items", actions, "", "No action items found.", " actions"),
        (c2, "Key decisions", decisions, "D", "No key decisions found.", ""),
        (c3, "Open questions", questions, "Q", "No open questions found.", ""),
    ]
    for col, name, entries, prefix, empty, extra in cols:
        with col:
            st.markdown(
                H(f'<div class="insight-group{extra}"><div class="insight-heading"><span>{name}</span>'
                  f'<b>{len(entries)}</b></div>{insight_items_html(entries, prefix, empty)}</div>'),
                unsafe_allow_html=True,
            )


def render_transcript_section(result: dict):
    transcript = result.get("transcript", "") or ""
    rows = transcript_rows(transcript)
    st.markdown(
        '<div class="section-heading"><div><span class="section-kicker">Recording</span>'
        '<div class="section-title">Transcript</div></div></div>',
        unsafe_allow_html=True,
    )
    t1, t2, t3 = st.columns([3, 1, 1])
    with t1:
        query = st.text_input("Search transcript", placeholder="Search transcript", label_visibility="collapsed")
    with t2:
        st.download_button("Download .txt", data=transcript, file_name="transcript.txt", mime="text/plain",
                           use_container_width=True, key="dl_txt_main")
    with t3:
        copy_button(transcript, "Copy transcript", key="copy_transcript")

    q = query.strip()
    shown = [(i, r) for i, r in enumerate(rows, 1) if not q or q.lower() in r.lower()]
    words = len(transcript.split())
    status = f"<span>{len(rows)} passages</span><span>{words:,} words</span>"
    if q:
        status += f"<span>{len(shown)} matching “{esc(q)}”</span>"
    st.markdown(f'<div class="transcript-status">{status}</div>', unsafe_allow_html=True)

    if not shown:
        st.markdown('<div class="empty-search">No passages match your search.</div>', unsafe_allow_html=True)
        return
    pattern = re.compile(re.escape(q), re.IGNORECASE) if q else None
    parts = []
    for i, r in shown:
        text = esc(r)
        if pattern:
            text = re.sub(pattern, lambda m: f"<mark>{m.group(0)}</mark>", text)
        parts.append(f'<div class="transcript-row"><span class="timestamp">{i:02d}</span><p>{text}</p></div>')
    with st.container(height=520, key="transcript-box"):
        st.markdown("".join(parts), unsafe_allow_html=True)


def ask(result: dict, question: str):
    st.session_state.chat_history.append({"role": "user", "content": question})
    with st.spinner("Thinking…"):
        try:
            answer = ask_question(result["rag_chain"], question)
        except Exception as e:  # noqa: BLE001
            answer = f"Something went wrong answering that: {e}"
    st.session_state.chat_history.append({"role": "assistant", "content": answer})


def render_chat_panel(result: dict):
    with st.container(border=True, key="chat-panel"):
        h1, h2 = st.columns([4, 1])
        with h1:
            st.markdown(
                '<div class="chat-head"><div class="section-title">Ask your meeting</div>'
                '<p>Search this transcript using natural language.</p></div>',
                unsafe_allow_html=True,
            )
        with h2:
            if st.session_state.chat_history and st.button("↺", key="chat_clear", help="Clear conversation"):
                st.session_state.chat_history = []
                st.rerun()
        st.markdown(
            H(f'<div class="chat-context"><i class="status-dot"></i>Grounded in this transcript'
              f'<span class="grow">{len(st.session_state.chat_history) // 2} questions</span></div>'),
            unsafe_allow_html=True,
        )
        with st.container(height=380, border=False, key="chat-scroll"):
            if not st.session_state.chat_history:
                st.markdown(
                    '<div class="chat-empty"><b>Ask about this meeting</b>'
                    '<p>Answers are retrieved from the transcript, not invented.</p></div>',
                    unsafe_allow_html=True,
                )
                for i, q in enumerate(SUGGESTED_QUESTIONS):
                    if st.button(q, key=f"sugg-{i}", use_container_width=True):
                        ask(result, q)
                        st.rerun()
            else:
                for i, msg in enumerate(st.session_state.chat_history):
                    with st.container(key=f"msg-{i}"):
                        if msg["role"] == "user":
                            st.markdown(
                                f'<span class="message-author user">You</span><p class="chat-user">{esc(msg["content"])}</p>',
                                unsafe_allow_html=True,
                            )
                        else:
                            st.markdown(f'<span class="message-author">{BRAND}</span>', unsafe_allow_html=True)
                            st.markdown(msg["content"])
        question = st.chat_input("Ask a question about this meeting…")
        st.markdown('<div class="chat-disclaimer">Answers are generated from the transcript and may need verification.</div>',
                    unsafe_allow_html=True)
        if question and question.strip():
            ask(result, question.strip())
            st.rerun()


def render_analysis_page(result: dict):
    render_meeting_bar(result)
    main, chat = st.columns([2.15, 1], gap="large")
    with main:
        render_summary_section(result)
        st.markdown('<div class="meeting-rule"></div>', unsafe_allow_html=True)
        render_insights_section(result)
        st.markdown('<div class="meeting-rule"></div>', unsafe_allow_html=True)
        render_transcript_section(result)
    with chat:
        render_chat_panel(result)


# ─────────────────────────────────────────────────────────────────────────
# Page: Saved & history
# ─────────────────────────────────────────────────────────────────────────
def render_history_page():
    st.markdown(
        H('<span class="context-label">Saved &amp; history</span><div class="page-title">Past analyses</div>'
          '<div class="intro-sub" style="margin:8px 0 22px;font-size:13px">Analyses from this browser session. '
          'They are cleared when the page is refreshed.</div>'),
        unsafe_allow_html=True,
    )
    items = list(reversed(st.session_state.history))
    if not items:
        st.markdown('<div class="empty-search">Nothing here yet — run an analysis and it will appear in this list.</div>',
                    unsafe_allow_html=True)
        if st.button("＋  New analysis", type="primary"):
            go("new")
            st.rerun()
        return
    q = st.text_input("Search analyses", placeholder="Search analyses", label_visibility="collapsed")
    items = [it for it in items if not q.strip() or q.lower() in it["title"].lower()]
    h = st.columns([4, 2, 1.3, 1.3])
    for col, label in zip(h, ("Name", "Analysed", "Length", "")):
        col.markdown(f'<div class="hist-head">{label}</div>', unsafe_allow_html=True)
    for i, it in enumerate(items):
        c = st.columns([4, 2, 1.3, 1.3])
        c[0].markdown(f'<div class="hist-cell hist-title">{esc(it["title"])}</div>', unsafe_allow_html=True)
        c[1].markdown(f'<div class="hist-cell">{esc(it.get("analyzed_at", ""))}</div>', unsafe_allow_html=True)
        c[2].markdown(f'<div class="hist-cell">{format_duration(it.get("duration_seconds"))}</div>', unsafe_allow_html=True)
        if c[3].button("Open", key=f"hist_open_{i}", use_container_width=True):
            open_result(it)
            st.rerun()


# ─────────────────────────────────────────────────────────────────────────
# Page: Settings
# ─────────────────────────────────────────────────────────────────────────
def pill(ok: bool, yes: str, no: str, optional: bool = False) -> str:
    if ok:
        return f'<span class="pill ok">{yes}</span>'
    return f'<span class="pill{"" if optional else " warn"}">{no}</span>'


def render_settings_page():
    st.markdown(
        H('<span class="context-label">Settings</span><div class="page-title">Workspace settings</div>'
          '<div class="intro-sub" style="margin:8px 0 10px;font-size:13px">Keys are read from your environment or '
          '.env file. Nothing is stored by this app.</div>'),
        unsafe_allow_html=True,
    )
    left, _ = st.columns([3, 1])
    with left:
        st.markdown('<div class="settings-title">Connections</div>', unsafe_allow_html=True)
        st.markdown(
            H(f'''<div class="setting-row"><div><b>Groq API key</b><p>LLM calls and hosted Whisper</p></div>
            {pill(bool(os.getenv("GROQ_API_KEY")), "Detected", "Missing")}</div>
            <div class="setting-row"><div><b>Sarvam API key</b><p>Only needed for Hinglish</p></div>
            {pill(bool(os.getenv("SARVAM_API_KEY")), "Detected", "Not set", optional=True)}</div>
            <div class="setting-row"><div><b>YouTube cookies</b><p>Helps with the “not a bot” check</p></div>
            {pill(any(os.getenv(k) for k in ("YT_COOKIES_FILE", "YT_COOKIES_B64", "YT_COOKIES_BROWSER")), "Configured", "Not set", optional=True)}</div>'''),
            unsafe_allow_html=True,
        )
        st.markdown('<div class="settings-title">Defaults</div>', unsafe_allow_html=True)
        st.markdown('<div class="field-label" style="margin-top:0">Transcript language</div>', unsafe_allow_html=True)
        with st.container(key="lang"):
            choice = st.radio("Default language", ["English", "Hinglish"], horizontal=True,
                              label_visibility="collapsed",
                              index=0 if st.session_state.default_language == "English" else 1, key="set_lang")
        st.session_state.default_language = choice

        st.markdown('<div class="settings-title danger-title">Session data</div>', unsafe_allow_html=True)
        st.caption(f"{len(st.session_state.history)} saved in this session.")
        if st.button("Clear session history", key="clear_hist"):
            st.session_state.history = []
            reset_to_input()
            st.session_state.page = "settings"
            st.rerun()


# ─────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────
inject_css()

_pending = st.session_state.pop("pending", None)
if _pending:
    st.session_state.page = "processing"

_page = st.session_state.page
if _page == "analysis" and not st.session_state.result:
    _page = st.session_state.page = "new"

render_sidebar()
render_top_header("new" if _page == "processing" else _page)

if _pending:
    run_pipeline(_pending["source"], _pending["language"], _pending.get("label") or _pending["source"])
elif _page == "analysis":
    render_analysis_page(st.session_state.result)
elif _page == "history":
    render_history_page()
elif _page == "settings":
    render_settings_page()
else:
    render_new_page()
